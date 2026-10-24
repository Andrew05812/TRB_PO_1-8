# -*- coding: utf-8 -*-
"""Общие функции оформления отчётов .docx для практик 3 и 4."""

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import os
import subprocess


SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"


def save_docx_and_pdf(doc, docx_path):
    """Сохраняет docx и конвертирует его в PDF рядом с исходником."""
    doc.save(docx_path)
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "pdf",
         "--outdir", os.path.dirname(docx_path), docx_path],
        check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
    )
    return os.path.splitext(docx_path)[0] + ".pdf"


def _set_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.orientation = WD_ORIENT.PORTRAIT

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.35
    style.paragraph_format.space_after = Pt(0)
    return doc


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    _set_font(p.add_run(text), size=15, bold=True)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _set_font(p.add_run(text), size=14 if level == 1 else 12, bold=True)


def add_paragraph(doc, text, indent=True, center=False, left=False):
    """Абзац отчёта.

    Кириллическую прозу выравниваем по ширине, а строки с командами - по левому
    краю. Иначе justify растягивает одиночные пробелы в команде на полстроки, и
    "dotnet publish -c Release" выглядит как таблица. Правило простое: если в
    тексте нет ни одной кириллической буквы, считаем его техническим.
    Явный left=True перекрывает автоопределение.
    """
    looks_like_command = left or bool(text) and not any(
        "а" <= ch <= "я" or ch == "ё" for ch in text.lower())
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif looks_like_command:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    if indent and not center and not looks_like_command:
        p.paragraph_format.first_line_indent = Cm(1.25)
    _set_font(p.add_run(text))


def _shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(paragraph.add_run(str(text)), size=10, bold=True)
        _shade(cell, "DCE6F1")

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _set_font(paragraph.add_run(str(value).replace("\n", " ")), size=10)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

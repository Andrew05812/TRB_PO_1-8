# -*- coding: utf-8 -*-
"""Собирает Отчет_Практика6.docx по результатам gdb и запусков."""

import os
import sys

from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "_tools"))
from docx_helper import add_heading, add_paragraph, add_table, add_title, new_document


def add_code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = 0
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    gdb = open(os.path.join(base, "results", "gdb_session.log"), encoding="utf-8").read()
    runs = open(os.path.join(base, "results", "runs_basic.txt"), encoding="utf-8").read()
    write = open(os.path.join(base, "results", "fmtwrite_write.txt"), encoding="utf-8").read()

    doc = new_document()

    add_title(doc, "Отчёт по практике 6")
    add_title(doc, "Бинарные уязвимости: format string, вариант 3")
    add_paragraph(doc, "", indent=False)
    add_paragraph(doc, "Выполнил: Колесников Андрей Юрьевич, группа БСБО-03-23, "
                       "студенческий билет 23Б0099", center=True, indent=False)
    add_paragraph(doc, "РТУ МИРЭА, Институт кибербезопасности и цифровых технологий",
                  center=True, indent=False)

    add_heading(doc, "1. Вариант и цель")
    add_paragraph(doc,
        "Номер студенческого билета 23Б0099 заканчивается на 9, поэтому по таблице "
        "вариантов взят вариант 3. В нём программа читает файл в буфер, а затем "
        "передаёт argv[1] напрямую в printf(). Требовалось провести динамическую "
        "отладку, разобрать код и предложить исправление. Уязвимость здесь не "
        "переполнение буфера, а управляемая строка формата (CWE-134), но она "
        "относится к тому же классу бинарных ошибок: пользовательский ввод меняет "
        "поведение машинного кода.")

    add_heading(doc, "2. Исходный код и сборка")
    add_paragraph(doc,
        "Исходный файл src/vuln_variant3.c взят из методички. Для наглядности в "
        "отладчике он собран 32-битным, без PIE и без stack protector:")
    add_code(doc,
        "gcc-12 -m32 -g -O0 -fno-stack-protector -no-pie -z execstack \\\n"
        "    -o build/vuln src/vuln_variant3.c")
    add_paragraph(doc,
        "Ключевой фрагмент:")
    add_code(doc,
        "FILE *secret = fopen(\"/challenge/app-systeme/ch5/.passwd\", \"rt\");\n"
        "char buffer[32];\n"
        "fgets(buffer, sizeof(buffer), secret);\n"
        "printf(argv[1]);          // ошибка: ввод стал строкой формата\n"
        "fclose(secret);")
    add_paragraph(doc,
        "Файл с «секретом» создаётся в контейнере и содержит строку "
        "SECRET{format_string_lab}. Проверок fopen и fgets в исходнике нет — это "
        "отдельные дефекты, но главная причина уязвимости в printf.")

    add_heading(doc, "3. Динамическая отладка")
    add_paragraph(doc,
        "В gdb поставлена точка останова на main+92 — непосредственно перед call "
        "printf, когда аргумент уже лежит на вершине стека. На первом запуске "
        "передан обычный текст AAAA. Видно, что буфер находится по адресу ebp-0x3c "
        "и содержит SECRET{format_string_lab}, а указатель FILE* лежит по ebp-0x1c. "
        "Строка формата — это argv[1], то есть данные, которые пользователь подал "
        "на вход программы.")
    add_code(doc, gdb[:2600])
    add_paragraph(doc,
        "Второй запуск с %p.%p... показывает, что printf читает аргументы дальше "
        "по стеку. Среди значений появляются 0x52434553 и 0x667b5445 — это байты "
        "SECR и ET{f из файла. Программа не собиралась печатать секрет, но он "
        "оказался в выводе из-за форматной строки. Это уже утечка памяти, а не "
        "безобидное эхо.")

    add_heading(doc, "4. Демонстрация примитивов")
    add_paragraph(doc,
        "Обычный запуск и утечка через %p:")
    add_code(doc, runs)
    add_paragraph(doc,
        "Для записи сделан отдельный стенд format_write_demo.c: он читает строку в "
        "стековый буфер и вызывает printf() от неё. В начале строки лежит адрес "
        "глобальной переменной target (0x0804c014), а %4$n пишет в неё число уже "
        "выведенных символов. Подбор смещения и результаты:")
    add_code(doc, write)
    add_paragraph(doc,
        "Так %62c%4$n записывает в target значение 0x42, а %66c%4$n — 0x46. "
        "Примитив произвольной записи доказан. В реальной программе вместо target "
        "можно выбрать адрес возврата или элемент GOT, если позволяют защиты; "
        "конкретный эксплойт зависит от сборки libc и включённых механизмов.")

    add_heading(doc, "5. Код-ревью")
    add_table(doc, ["Проблема", "Где", "CWE", "Последствие"], [
        ["Ввод передан как строка формата", "printf(argv[1])", "CWE-134",
         "чтение памяти, запись через %n, падение"],
        ["Нет проверки argc", "main", "CWE-20",
         "обращение к argv[1] при пустом argv"],
        ["Нет проверки fopen", "fopen(...)", "CWE-252",
         "fgets(NULL) при отсутствии файла"],
        ["Секрет хранится в памяти процесса", "buffer[32]", "CWE-316",
         "утечка через %s/%p или дамп"],
        ["Отключены защиты", "сборка -no-pie -z execstack", "CWE-693",
         "упрощает эксплуатацию, но это сделано для лабораторной"],
    ])
    add_paragraph(doc,
        "Главная ошибка — смешение данных и управляющей строки. printf получает "
        "первый аргумент как программу форматирования: %p читает слово со стека, "
        "%s разыменовывает указатель, %n записывает число по адресу. Даже если "
        "программа не печатает результат, %n может изменить память и привести к "
        "выполнению чужого кода. В варианте 3 есть все условия: управляемый первый "
        "аргумент, отсутствие проверки формата и полезные данные в памяти.")

    add_heading(doc, "6. Исправление")
    add_paragraph(doc,
        "Исправление минимальное: строка формата должна быть константой, а ввод — "
        "аргументом:")
    add_code(doc, "printf(\"%s\\n\", argv[1]);")
    add_paragraph(doc,
        "В fixed_variant3.c дополнительно проверяются argc, fopen и fgets. Это не "
        "лечит форматную строку (её лечит только %s), но убирает падения при "
        "неправильном запуске. Проверка исправленной версии:")
    add_code(doc,
        "$ ./build/fixed \"ARGV=%p argv2=%p %n\"\n"
        "ARGV=%p argv2=%p %n")
    add_paragraph(doc,
        "Символы % больше не выполняются. Для реальной сборки нужно также включить "
        "-Wformat-security, -D_FORTIFY_SOURCE=2, PIE, stack protector и не "
        "отключать ASLR. Это не заменяет исправление, но повышает цену эксплуатации "
        "похожих ошибок.")

    add_heading(doc, "7. Выводы")
    add_paragraph(doc,
        "В варианте 3 подтверждена format string vulnerability. В gdb видно, что "
        "printf получает пользовательский указатель как формат; практические запуски "
        "показывают утечку SECRET из буфера и запись 0x42/0x46 в глобальную "
        "переменную через %n. Исправление на одну строку (printf(\"%s\", argv[1])) "
        "полностью убирает управление форматом, а дополнительные проверки argc, "
        "fopen и fgets убирают аварийные ситуации.")
    add_paragraph(doc,
        "Ограничение: ASLR в контейнере полностью отключить не удалось, адреса стека "
        "меняются. На демонстрацию это не влияет, потому что смещение %4$n и адрес "
        "target фиксированы (сборка -no-pie). Все действия выполнялись в локальном "
        "учебном контейнере, файл с секретом — тестовый.")

    out = os.path.join(base, "Отчет_Практика6.docx")
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
